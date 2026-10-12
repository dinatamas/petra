#!/usr/bin/env python3
from copy import copy
from pathlib import Path
from typing import NamedTuple
import re
import sys

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd


EXCEL_DIR = Path("./data/")
RESULT_FILE = "macskas.xlsx"
HIDE_UNUSED_CELLS = False
PRINT_BEHAVIOR_NAMES = False
BEHAVIOR_NAMES = set()
DEBUG = False


class Behavior(NamedTuple):
    name: str
    time: float


class Trial(NamedTuple):
    valid: bool
    valid_count: int
    lat_pass: float
    lat_arrive: float
    lat_hesit_start: float
    lat_hesit_pass: float
    freq_attempt_pass: int
    freq_alternative_side: int
    freq_lookback_exp: int
    freq_look_up: int
    freq_towards_stops: int
    freq_approach_again: int
    freq_walk_away: int
    freq_gave_up: int
    freq_pass_invalid: int


def main():
    columns = ("name", *Trial._fields)
    horizontal_rows, vertical_rows = [], []

    for subdir in sorted(EXCEL_DIR.iterdir()):
        if not subdir.is_dir():
            continue

        try:
            horizontal, vertical = load_files(subdir)
        except DataLoadError, InvalidDataError:
            continue

        horizontal_rows.extend(horizontal)
        vertical_rows.extend(vertical)

    if PRINT_BEHAVIOR_NAMES:
        print("\n [i] Behavior oszlop értékei:\n")
        for name in sorted(BEHAVIOR_NAMES):
            print(f"    - {name}")
        sys.exit()

    horizontal_df = pd.DataFrame(horizontal_rows, columns=columns)
    vertical_df = pd.DataFrame(vertical_rows, columns=columns)

    with pd.ExcelWriter(RESULT_FILE, engine="openpyxl") as writer:

        horizontal_df.to_excel(writer, sheet_name="Horizontal", index=False)
        vertical_df.to_excel(writer, sheet_name="Vertical", index=False)

        # Styles used in the worksheet.
        header_fill = PatternFill(fill_type="solid", fgColor="000000")
        header_font = Font(bold=True, color="FFFFFF")
        light_green_fill = PatternFill(fill_type="solid", bgColor="FFC6EFCE")
        dark_green_fill = PatternFill(fill_type="solid", bgColor="FF81C784")
        light_red_fill = PatternFill(fill_type="solid", bgColor="FFFFC7CE")
        dark_red_fill = PatternFill(fill_type="solid", bgColor="FFE57373")
        gray_side = Side(style="thick", color="FF404040")

        for worksheet in writer.sheets.values():

            # Expand each column to fit the longest text in it.
            for cells in worksheet.columns:
                max_length = max(len(str(c.value)) if c.value is not None else 0 for c in cells)
                column_letter = get_column_letter(cells[0].column)
                worksheet.column_dimensions[column_letter].width = max_length + 1

            # Make first row and column inverse colored and bold.
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
            for row in worksheet.iter_rows(min_col=1, max_col=1):
                row[0].fill = header_fill
                row[0].font = header_font

            # Hide unpopulated rows and columns at the worksheet edges.
            if HIDE_UNUSED_CELLS:
                hide_row_start, hide_row_end = worksheet.max_row + 1, 1048576
                hide_col_start, hide_col_end = get_column_letter(worksheet.max_column + 1), get_column_letter(16384)
                worksheet.row_dimensions.group(hide_row_start, hide_row_end, hidden=True)
                worksheet.column_dimensions.group(hide_col_start, hide_col_end, hidden=True)

            # Conditionally color frequencies when data is available.
            data_range = f"H2:{worksheet.cell(row=worksheet.max_row, column=worksheet.max_column).coordinate}"
            worksheet.conditional_formatting.add(data_range, FormulaRule(formula=['=AND($B2=TRUE,H2>0)'], fill=dark_green_fill))
            worksheet.conditional_formatting.add(data_range, FormulaRule(formula=['=AND($B2=FALSE,H2>0)'], fill=dark_red_fill))

            # Conditionally color based on attempt success.
            data_range = f"B2:{worksheet.cell(row=worksheet.max_row, column=worksheet.max_column).coordinate}"
            worksheet.conditional_formatting.add(data_range, FormulaRule(formula=['$B2=TRUE'], fill=light_green_fill))
            worksheet.conditional_formatting.add(data_range, FormulaRule(formula=['$B2=FALSE'], fill=light_red_fill))

            # Add an outer border around the entire populated part of the sheet.
            for row in range(1, worksheet.max_row + 1):
                for col in range(1, worksheet.max_column + 1):
                    cell = worksheet.cell(row=row, column=col)
                    border = copy(cell.border)
                    if row == 1:
                        border.top = gray_side
                    if row == worksheet.max_row:
                        border.bottom = gray_side
                    if col == 1:
                        border.left = gray_side
                    if col == worksheet.max_column:
                        border.right = gray_side
                    cell.border = border

            # Add an inner border when the subject changes.
            for row in range(1, worksheet.max_row):
                current_name = worksheet.cell(row=row, column=1).value
                next_name = worksheet.cell(row=row + 1, column=1).value
                if current_name != next_name:
                    for col in range(1, worksheet.max_column + 1):
                        cell = worksheet.cell(row=row, column=col)
                        border = copy(cell.border)
                        border.bottom = gray_side
                        cell.border = border

    green("A macskas.xlsx file sikeresen elkészítve!")


def load_files(subdir):
    files = [f for f in subdir.iterdir() if f.suffix.lower() in (".xls", ".xlsx") and not f.name.startswith("~$")]
    if len(files) != 2:
        red(f" [!] A(z) {subdir.name} mappában {len(files)} Excel fájl van, nem 2.")
        raise DataLoadError

    def find_file_with_word(word):
        matches = [f for f in files if word in f.name.lower()]
        if len(matches) == 0:
            red(f' [!] A(z) {subdir.name} mappában nincs "{word}" Excel file.')
            raise DataLoadError
        if len(matches) > 1:
            red(f' [!] A(z) {subdir.name} mappában túl sok file nevében szerepel a "{word}" szó.')
            raise DataLoadError
        return matches[0]

    horizontal = find_file_with_word("horizontalis")
    horizontal_trials = process_excel(horizontal)
    horizontal_rows = [(subdir.name, *trial) for trial in horizontal_trials]

    vertical = find_file_with_word("vertikalis")
    vertical_trials = process_excel(vertical)
    vertical_rows = [(subdir.name, *trial) for trial in vertical_trials]

    return horizontal_rows, vertical_rows


def process_excel(file):
    df = pd.read_excel(file, usecols=["Behavior", "Start (s)"])
    behaviors = [Behavior(str(row["Behavior"]), float(row["Start (s)"])) for _, row in df.iterrows()]
    starts = [i for i, b in enumerate(behaviors) if re.fullmatch(r"start\d+", b.name)]
    groups = [behaviors[start:end] for start, end in zip(starts, starts[1:] + [len(behaviors)])]
    trials = []

    for behavior in behaviors:
        BEHAVIOR_NAMES.add(behavior.name)

    for i, group in enumerate(groups):

        def find_behaviors(regex, min=None, max=None):
            matches = [behavior for behavior in group if re.fullmatch(regex, behavior.name)]
            if min and len(matches) < min:
                red(f" [!] A {file} file {i+1}. tesztben {len(matches)} db {regex} van (min: {min})")
                raise InvalidDataError
            if max and len(matches) > max:
                red(f" [!] A {file} file {i+1}. tesztben {len(matches)} db {regex} van (max: {max})")
                raise InvalidDataError
            if max == 1:
                return matches[0] if matches else None
            return matches

        start = find_behaviors(r"start\d+", min=1, max=1)
        stop = find_behaviors(r"stop\d+", max=1)
        approach_start = find_behaviors(r"approach_start", min=1, max=1)
        arrive_door = find_behaviors(r"arrive_door", max=1)
        gave_up = find_behaviors(r"gave_up", max=1)
        walk_away = find_behaviors(r"walk_away")
        approach_again = find_behaviors(r"approach_again")
        pass_invalid = find_behaviors(r"pass_invalid", max=1)

        valid = stop is not None

        # Gave up is always failed.
        if gave_up:
            if stop:
                cyan(f' [x] A {file} file {i+1}. tesztben van gave_up és stop is.')
            elif group[-1].name != "gave_up":
                red(f' [!]  A {file} file {i+1}. tesztben a gave_up után van más behavior.')
            valid = False

        # If walked away without reapproaching, it is also invalid.
        if walk_away and (not approach_again or approach_again[-1].time <= walk_away[0].time):
            if stop:
                cyan(f' [x] A {file} file {i+1}. tesztben van walk_away aztkövető approach_again nélkül és stop is.')
            elif group[-1].name != "walk_away":
                red(f' [!]  A {file} file {i+1}. tesztben a walk_away után van más behavior.')
            valid = False

        # Invalid pass attempts also fail the trial.
        if pass_invalid:
            if stop:
                cyan(f' [x] A {file} file {i+1}. tesztben van pass_invalid és stop is.')
            elif group[-1].name != "pass_invalid":
                red(f' [!]  A {file} file {i+1}. tesztben a pass_invalid után van más behavior.')
            valid = False

        # The trial cannot pass without arriving at the door.
        if not arrive_door:
            if stop:
                cyan(f' [x] A {file} file {i+1}. tesztben van stop de nincs arrive_door.')
            valid = False

        if stop and group[-1].name[:4] != "stop" and stop.time > group[-1].time:
            yellow(f' [?] A {file} file {i+1}. tesztben van behavior a stop után.')

        trials.append(Trial(
            valid = valid,
            valid_count = 0,  # Temporarily.
            lat_pass = stop.time - approach_start.time if valid else None,
            lat_arrive = arrive_door.time - approach_start.time if arrive_door else None,
            lat_hesit_start = approach_start.time - start.time,
            lat_hesit_pass = stop.time - arrive_door.time if valid and arrive_door else None,
            freq_attempt_pass = len(find_behaviors(r"attempt_pass")),
            freq_alternative_side = len(find_behaviors(r"alternative_side")),
            freq_lookback_exp = len(find_behaviors(r"lookback_exp")),
            freq_look_up = len(find_behaviors(r"look_up")),
            freq_towards_stops = len(find_behaviors(r"towards_stops")),
            freq_approach_again = len(find_behaviors(r"approach_again")),
            freq_walk_away = len(find_behaviors(r"walk_away")),
            freq_gave_up = len(find_behaviors(r"gave_up")),
            freq_pass_invalid = len(find_behaviors(r"pass_invalid")),
        ))

        # If the attempt did not pass, do not consider future trials.
        if not valid:
            if i != len(groups) - 1:
                cyan(f' [!] A {file} file {i+1}. teszt bukott de van következő trial.')
            break

    valid_count = len([t for t in trials if t.valid])
    trials = [Trial(**{**t._asdict(), "valid_count": valid_count}) for t in trials]

    return trials


class InvalidDataError(Exception):
    pass


class DataLoadError(Exception):
    pass


def red(msg):
    print(f"\n\033[31m{msg}\033[0m")

def green(msg):
    print(f"\n\033[32m{msg}\033[0m")

def yellow(msg):
    print(f"\n\033[33m{msg}\033[0m")

def cyan(msg):
    if DEBUG:
        print(f"\n\033[36m{msg}\033[0m")


if __name__ == "__main__":
    main()
